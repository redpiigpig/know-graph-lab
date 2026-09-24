# -*- coding: utf-8 -*-
"""一貫道國家檔案分類總表：catalog/*.json → 加密 Excel。

    python -X utf8 scripts/yiguandao_catalog_xlsx.py            # 全部卷宗各出一個分頁
    python -X utf8 scripts/yiguandao_catalog_xlsx.py --cases BS  # 只出指定卷宗的分頁

正本是 public/content/research-data/yiguandao/catalog/ 的 cases.json 與 items.json
（整夾 gitignore：裡面有人名）。Excel 是從正本產生的成品，各卷宗分頁每次重產；
要改資料就改 items.json 或 Excel 總表後再匯回，別只改某個卷宗分頁。

一列＝一件公文，不是一頁：卷宗會到幾百個，只有「件」這個單位能跨卷宗篩選。
全文不在這裡——全文正本是人工謄打的 Word，總表只記定位（PDF 頁／原檔頁）。

🚨 輸出用 .env 的 YIGUANDAO_DOCX_PASSWORD 加密（與檔案全文 Word 同一組）。
   未加密的中間檔只落在暫存目錄，寫完即刪。
"""
import argparse
import io
import json
import os
import tempfile
from collections import defaultdict
from pathlib import Path

import msoffcrypto  # noqa: F401  (OOXMLFile 需要)
from msoffcrypto.format.ooxml import OOXMLFile
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
CAT = ROOT / "public/content/research-data/yiguandao/catalog"
DEST = Path(r"G:/我的雲端硬碟/玄奘/博一上/工作/中研院助理/08_一貫道國家檔案分類總表.xlsx")

# 分類詞表：單值欄做下拉選單，多值欄（以「；」分隔）只列建議值
VOCAB = {
    "時期": ["抗戰", "戰後", "戡亂", "遷台初期", "戒嚴", "解嚴後"],
    "機關類型": ["軍事", "情治", "行政", "司法", "黨務", "地方政府", "民間團體", "報刊"],
    "行動": ["建檔", "發動查禁", "通令查禁", "覆文", "偵查", "情報", "情報分送", "逮捕", "撤壇",
             "查封", "判決", "不偵查", "移司法", "行政處理", "陳情", "訴願", "駁回", "指控", "檢舉",
             "詢問", "請示", "調卷", "研判", "研議", "擱置", "補報", "轉呈", "交辦", "批覆",
             "備查", "送達", "覆查", "爭議", "產權糾紛", "開禁"],
    "定性框架": ["敵偽", "附匪", "奸黨嫌疑", "邪教", "迷信", "邪說", "斂財", "秘密結社",
                 "軍人入道", "易為奸偽利用"],
    "文種": ["公函", "訓令", "指令", "代電", "代電稿", "極密代電", "快郵代電", "電", "呈", "簽呈",
             "簽註", "批", "判決", "訴願決定書", "報告", "調查報告", "情報", "名冊", "清冊", "表",
             "剪報", "章程", "信封", "封面", "目次表", "摘要卡"],
}
SINGLE = ["時期", "機關類型", "行動"]      # 這三欄一件只填一個值，可做下拉
NOTES = {
    "時期": "抗戰＝至34年8月；戰後＝34年8月至36年7月；戡亂＝36年7月動員戡亂令起",
    "定性框架": "公文對一貫道的指稱或指控，照公文用語歸類；可多值，以「；」分隔",
    "行動": "該件公文本身做的事，不是整案的結果",
}

HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")
BOLD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
WIDTH = {"卷宗代號": 8, "件序": 5, "PDF頁": 8, "原檔頁": 12, "日期原文": 14, "日期（國曆）": 11,
         "發文者": 20, "受文者": 18, "文種": 10, "字號": 18, "事由摘要": 48, "時期": 7,
         "機關類型": 9, "行動": 9, "定性框架": 12, "處置結果": 12, "地點": 12, "人名": 24, "備註": 22}


def header(ws, cols, widths=None):
    ws.append(cols)
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=1, column=i)
        cell.font, cell.fill = BOLD, HEAD_FILL
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(c, 16)
    ws.freeze_panes = "A2"


def item_sheet(ws, cols, rows):
    header(ws, cols, WIDTH)
    for r in rows:
        ws.append([r.get(c, "") for c in cols])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
    ws.auto_filter.ref = ws.dimensions


VOCAB_COL = {}


def build(case_filter=None):
    cases = json.loads((CAT / "cases.json").read_text(encoding="utf-8"))
    data = json.loads((CAT / "items.json").read_text(encoding="utf-8"))
    cols, items = data["columns"], data["items"]
    people_info = {}
    pf = CAT / "people.json"
    if pf.exists():
        people_info = json.loads(pf.read_text(encoding="utf-8"))

    wb = Workbook()
    ws = wb.active
    ws.title = "說明"
    ws.column_dimensions["A"].width = 110
    for line in [
        "一貫道國家檔案分類總表",
        "",
        "結構",
        "・卷宗索引：一列＝一卷宗。",
        "・總表：一列＝一件公文，所有卷宗都在這一張，跨卷宗查詢與篩選都在這裡做（這是正本）。",
        "・各卷宗分頁（以卷宗代號命名）：由總表自動產生，只供逐卷閱讀；修改請改總表。",
        "・分類詞表：時期、機關類型、行動三欄可從下拉選單選；定性框架可多值，以「；」分隔。",
        "・人名權威檔：一人一列，列出身分與出現的件。",
        "",
        "定位",
        "・PDF頁：謄打本 Word／PDF 的頁碼；原檔頁：檔案原件蓋印的頁碼。",
        "・全文不在本表；全文正本為人工謄打的 Word 檔（見卷宗索引「Word檔」欄）。",
        "",
        "判讀",
        "・日期原文照錄（含韻目代日），日期（國曆）為換算。",
        "・備註標「待核影像」者，引用前須回原件影像核對。",
        "・人名欄僅記公文中出現者，未經查證者不代表身分已確認。",
        "",
        "本檔含第三人姓名，限計畫內部使用，請勿外傳。",
    ]:
        ws.append([line])
    ws["A1"].font = Font(bold=True, size=14)
    for r in (3, 10, 14):
        ws.cell(row=r, column=1).font = BOLD

    ws = wb.create_sheet("卷宗索引")
    ccols = ["卷宗代號", "檔號", "案名", "典藏機關", "全宗", "起訖", "影像張數", "謄打本頁數",
             "謄打範圍", "公開狀態", "Word檔", "網址", "摘要"]
    header(ws, ccols, {"檔號": 26, "案名": 22, "起訖": 24, "謄打範圍": 30, "公開狀態": 20,
                       "Word檔": 26, "網址": 30, "摘要": 60})
    for c in cases:
        ws.append([c.get(k, "") for k in ccols])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP

    ws = wb.create_sheet("總表")
    item_sheet(ws, cols, items)

    by_case = defaultdict(list)
    for it in items:
        by_case[it["卷宗代號"]].append(it)
    wanted = case_filter or [c["卷宗代號"] for c in cases]
    names = {c["卷宗代號"]: c["案名"] for c in cases}
    for code in wanted:
        title = f"{code} {names.get(code, '')}"[:31].replace("/", "／")
        for ch in "[]:*?\\":
            title = title.replace(ch, "")
        item_sheet(wb.create_sheet(title), cols, by_case.get(code, []))

    ws = wb.create_sheet("分類詞表")
    vcols = list(VOCAB)
    header(ws, vcols)
    for i, name in enumerate(vcols, 1):
        VOCAB_COL[name] = get_column_letter(i)
        for j, v in enumerate(VOCAB[name], 2):
            ws.cell(row=j, column=i, value=v)
    r0 = max(len(v) for v in VOCAB.values()) + 3
    ws.cell(row=r0, column=1, value="欄位說明").font = BOLD
    for k, (name, note) in enumerate(NOTES.items(), 1):
        ws.cell(row=r0 + k, column=1, value=name)
        ws.cell(row=r0 + k, column=2, value=note)

    # 下拉選單要指向詞表的欄，所以詞表建好後再補上驗證
    for sheet in wb.worksheets:
        if sheet.title in ("說明", "卷宗索引", "分類詞表", "人名權威檔"):
            continue
        sheet.data_validations.dataValidation = []
        for name in SINGLE:
            col = get_column_letter(cols.index(name) + 1)
            dv = DataValidation(type="list", formula1=f"=分類詞表!${VOCAB_COL[name]}$2:${VOCAB_COL[name]}$60",
                                allow_blank=True, showErrorMessage=False)
            sheet.add_data_validation(dv)
            dv.add(f"{col}2:{col}5000")

    ws = wb.create_sheet("人名權威檔")
    pcols = ["人名", "其他寫法", "身分／職位", "派系或組線", "後來去向", "出現（卷宗-件序）", "出現件數", "查證出處"]
    header(ws, pcols, {"身分／職位": 34, "派系或組線": 20, "後來去向": 28, "出現（卷宗-件序）": 30, "查證出處": 24})
    seen = defaultdict(list)
    for it in items:
        for n in (it.get("人名") or "").split("；"):
            n = n.strip()
            if n:
                seen[n].append(f"{it['卷宗代號']}-{it['件序']}")
    for n in sorted(seen, key=lambda x: (-len(seen[x]), x)):
        info = people_info.get(n, {})
        ws.append([n, info.get("其他寫法", ""), info.get("身分", ""), info.get("派系", ""),
                   info.get("去向", ""), "、".join(seen[n]), len(seen[n]), info.get("出處", "")])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
    ws.auto_filter.ref = ws.dimensions
    return wb, len(items), len(seen)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--cases", help="只產生這些卷宗代號的分頁，逗號分隔")
    ap.add_argument("--out", default=str(DEST))
    a = ap.parse_args()
    pw = os.environ.get("YIGUANDAO_DOCX_PASSWORD")
    if not pw:
        for line in (ROOT / ".env").read_text(encoding="utf-8").splitlines():
            if line.startswith("YIGUANDAO_DOCX_PASSWORD="):
                pw = line.split("=", 1)[1].strip()
    if not pw:
        raise SystemExit("缺 YIGUANDAO_DOCX_PASSWORD")
    wb, n_items, n_people = build(a.cases.split(",") if a.cases else None)
    with tempfile.TemporaryDirectory() as td:
        plain = Path(td) / "plain.xlsx"
        wb.save(plain)
        buf = io.BytesIO()
        with open(plain, "rb") as f:
            OOXMLFile(f).encrypt(pw, buf)
    Path(a.out).write_bytes(buf.getvalue())
    print(f"{n_items} 件、{n_people} 人 → {a.out}（已加密）")


if __name__ == "__main__":
    main()
