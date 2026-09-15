# -*- coding: utf-8 -*-
"""把四門課的修課名單收成一份正規化 JSON，並寫進 Drive 各課程資料夾。

名單有三個來源，優先序由高而低：

1. **教職員服務系統匯出的「課堂管理名冊」`.xls`**（BIFF，OLE2）——唯一逐個學生
   都有的來源。系統匯出的檔名是 UUID，所以認檔靠讀內容的「科目代號:」而不是檔名。
2. **Drive 各課程資料夾的 `*-點名表.docx`** —— 由上面那批產的快照，`PPA001`
   目前只有這一份（那門是代課，沒拿到 xls）。
3. **公開開課查詢的「已選人數」** —— 只有數字，拿來對帳，不能當名單。
   🚨 三者不一致是常態：加退選結束前系統天天在動，而 xls／docx 是當天的快照。

🚨 名單含學生個資，**只寫 Drive 不進 git**（CLAUDE.md 第一條）。

用法：
    python scripts/course_roster.py            # 印對帳表，不寫檔
    python scripts/course_roster.py --write    # 一併寫進 Drive 各課程資料夾
"""
import argparse
import datetime as dt
import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / 'scripts'))

import course_schedule as CS  # noqa: E402

DRIVE = Path(r'G:\我的雲端硬碟\資料\知識圖工作室\教學')
DOWNLOADS = Path.home() / 'Downloads'

# 課號 → (Drive 資料夾, 點名表檔名)
FOLDER = {
    'BBE275': ('115-1_世界宗教文化導論', '世界宗教文化導論-點名表.docx'),
    'PPA001': ('115-1_世界宗教文化導論', '二年制世界宗教文化導論-點名表.docx'),
    'BBE150': ('115-1_基督宗教概論', '基督宗教概論-點名表.docx'),
    'PPA066': ('115-1_宗教系國文講義', '國文-點名表.docx'),
}

# 點名表／名冊姓名欄後面掛的學籍註記。表尾自己寫著：
# 休:休學 退:退學 刪:刪除 保:保留學籍 W:期中退選。這些人不會來上課。
MARK = re.compile(r'[（(](休學?|退學?|刪除?|保留?|W)[）)]')


def _clean(name):
    """把姓名欄拆成（乾淨姓名, 學籍註記或 None）。

    🚨 不是每個括號都是學籍註記：`詹淑禎(傳聞)` 那個括號是法名，要留著。
    """
    m = MARK.search(name)
    if not m:
        return name.strip(), None
    return MARK.sub('', name).strip(), m.group(1)


def from_xls(path):
    """讀一份「課堂管理名冊」xls。回傳 dict，讀不出科目代號就回 None。"""
    import xlrd
    sh = xlrd.open_workbook(str(path)).sheet_by_index(0)
    head = ' '.join(str(c.value) for r in range(min(4, sh.nrows)) for c in sh.row(r))
    m = re.search(r'科目代號[:：]\s*(\S+)', head)
    if not m:
        return None
    students = []
    for i in range(sh.nrows):
        c = [str(x.value).strip() for x in sh.row(i)]
        if len(c) > 3 and re.fullmatch(r'\d+(\.0)?', c[0]) and c[1]:
            nm, mark = _clean(c[2])
            students.append({'sid': c[1], 'name': nm, 'klass': c[3], 'mark': mark})
    return {
        'code': m.group(1),
        'source': 'xls',
        'source_file': path.name,
        'snapshot': dt.date.fromtimestamp(path.stat().st_mtime).isoformat(),
        'students': students,
    }


def from_docx(path):
    from docx import Document
    students = []
    for t in Document(str(path)).tables:
        for r in t.rows:
            c = [x.text.strip() for x in r.cells]
            if len(c) > 3 and re.fullmatch(r'\d+', c[0] or '') and c[1]:
                nm, mark = _clean(c[2])
                students.append({'sid': c[1], 'name': nm, 'klass': c[3], 'mark': mark})
    return {
        'source': 'docx',
        'source_file': path.name,
        'snapshot': dt.date.fromtimestamp(path.stat().st_mtime).isoformat(),
        'students': students,
    }


def collect():
    """回傳 {課號: 名單 dict}。"""
    out = {}
    for f in sorted(DOWNLOADS.glob('*.xls')):
        try:
            d = from_xls(f)
        except Exception:
            continue
        if d and d['code'] in FOLDER:
            # 同一門課有多份就取最新的快照
            if d['code'] not in out or d['snapshot'] >= out[d['code']]['snapshot']:
                out[d['code']] = d
    for code, (folder, docname) in FOLDER.items():
        if code in out:
            continue
        p = DRIVE / folder / docname
        if p.exists():
            d = from_docx(p)
            d['code'] = code
            out[code] = d
    return out


def payload(rosters, live):
    """組成網頁與存檔共用的那一份資料。"""
    courses = []
    for key, c in CS.COURSES.items():
        code = c['code']
        r = rosters.get(code, {'students': [], 'source': None,
                               'source_file': None, 'snapshot': None})
        # 🚨 休學／退學／刪除／保留學籍／期中退選的人**不進 students**。
        # 使用者：「休學的就不要再放上來了，我也無法給他們打分數」。
        # 他們只留在 excluded，供對帳交代 49 人怎麼變 40 人。
        active = [s for s in r['students'] if not s['mark']]
        excluded = [s for s in r['students'] if s['mark']]
        courses.append({
            'key': key,
            'code': code,
            'name': c['name'],
            'klass': c['klass'],
            'time': c['time'],
            'room': c['room'],
            'assessment': [{'item': i, 'pct': int(p.rstrip('%'))}
                           for i, p in c['assessment']],
            'source': r['source'],
            'source_file': r['source_file'],
            'snapshot': r['snapshot'],
            'listed': len(r['students']),
            'active': len(active),
            'system': live.get(code),
            'students': active,
            'excluded': excluded,
        })
    return {'year': '115', 'term': '1', 'teacher': CS.TEACHER,
            'generated': dt.date.today().isoformat(), 'courses': courses}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--write', action='store_true', help='寫進 Drive 各課程資料夾')
    ap.add_argument('--no-live', action='store_true', help='不連校網查已選人數')
    a = ap.parse_args()
    sys.stdout.reconfigure(encoding='utf-8')

    live = {}
    if not a.no_live:
        import course_enrollment as CE
        live = CE.fetch()

    data = payload(collect(), live)

    print(f'{data["year"]}-{data["term"]}　{data["teacher"]}　'
          f'（{data["generated"]} 查）\n')
    print(f'{"課號":<8}{"課程":<12}{"名單":>5}{"扣註記":>7}{"系統已選":>9}  來源')
    for c in data['courses']:
        flag = '' if c['system'] in (None, c['active']) else '  ⚠ 不一致'
        print(f'{c["code"]:<8}{c["name"]:<12}{c["listed"]:>5}{c["active"]:>7}'
              f'{str(c["system"]):>9}  {c["source"] or "—"}'
              f'（{c["snapshot"] or "—"}）{flag}')

    if not a.write:
        print('\n（未寫檔；要寫進 Drive 加 --write）')
        return

    for c in data['courses']:
        if not c['students']:
            print(f'略過 {c["code"]}：沒有名單來源')
            continue
        folder = DRIVE / FOLDER[c['code']][0]
        # 🚨 檔名用**來源快照日**不是今天：名單是那天匯出的，寫今天會讓人
        # 以為這份是今天的系統狀態（BBE275 就差了 8 個人）。
        stem = f'{c["code"]}_{c["name"]}_選課名單_{c["snapshot"]}'
        (folder / f'{stem}.json').write_text(
            json.dumps(c, ensure_ascii=False, indent=2), encoding='utf-8')
        write_xlsx(folder / f'{stem}.xlsx', data, c)
        kinds = '.json / .xlsx'
        # 系統匯出的原檔檔名是 UUID，留在下載夾等於丟了。連同原檔一起歸檔。
        if c['source'] == 'xls':
            src = DOWNLOADS / c['source_file']
            if src.exists():
                (folder / f'{stem}_課堂管理名冊.xls').write_bytes(src.read_bytes())
                kinds += ' / .xls（系統原檔）'
        print(f'寫入 {folder / stem}{kinds}')


def write_xlsx(path, data, c):
    """一門課一張表：名單＋評量項目空欄，直接可以拿來登分。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = c['code']
    head = Font(bold=True, color='FFFFFF')
    fill = PatternFill('solid', fgColor='1E5A4C')

    ws.append([f'{data["year"]}-{data["term"]}　{c["name"]}（{c["code"]}）'])
    ws.append([f'{c["klass"]}　{c["time"]}　{c["room"]}　授課：{data["teacher"]}'])
    ws.append([f'在籍可登分 {c["active"]} 人（原名單 {c["listed"]} 人、'
               f'扣學籍註記 {c["listed"] - c["active"]} 人）、系統已選 {c["system"]}　'
               f'來源 {c["source_file"]}（{c["snapshot"]} 快照）　產生於 {data["generated"]}'])
    if c['excluded']:
        ws.append(['不列入（學籍註記）：' + '、'.join(
            f'{s["name"]}（{s["mark"]}）' for s in c['excluded'])])
    ws.append([])

    cols = ['序', '學號', '姓名', '班級'] + \
           [f'{i["item"]} {i["pct"]}%' for i in c['assessment']] + ['總成績']
    ws.append(cols)
    r0 = ws.max_row
    for cell in ws[r0]:
        cell.font = head
        cell.fill = fill
        cell.alignment = Alignment(horizontal='center', wrap_text=True)

    first = 5  # 評量項目的第一欄（E）
    last = first + len(c['assessment']) - 1
    for n, s in enumerate(c['students'], 1):
        ws.append([n, s['sid'], s['name'], s['klass']])
        row = ws.max_row
        terms = '+'.join(
            f'{get_column_letter(first + i)}{row}*{it["pct"]}/100'
            for i, it in enumerate(c['assessment']))
        ws.cell(row, last + 1).value = f'=IF(COUNT({get_column_letter(first)}{row}:' \
                                       f'{get_column_letter(last)}{row})=0,"",ROUND({terms},0))'

    widths = [4, 12, 14, 14] + [11] * len(c['assessment']) + [9]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f'A{r0 + 1}'
    wb.save(str(path))


if __name__ == '__main__':
    main()
