#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""《神的演化》（羅伯‧賴特著、梁永安譯，大家出版 2013）OCR → 分章朗讀 docx。

使用者 2026-10-01 要的，格式比照 G:\\我的雲端硬碟\\神的歷史\\神的歷史（分章朗讀）：
章名＝標題 1、小節＝標題 2、正文一段一段；**註釋不收**（朗讀用）。

來源：掃描本 526 頁、直排、無文字層。OCR 主力是 Gemini（c:/tmp/evo/gem/*.json，
scan_ocr 的結構化行格式）；MinerU 在這本會**大量吃掉標點**（「……譴責當時拙著……」
整段沒有逗號句號），朗讀檔不能用，只在 Gemini 缺頁時拿 c:/tmp/evo/part_*.jsonl 頂。

分章靠目次的印刷頁碼；PDF 頁＝印刷頁＋6（序言印刷頁 4 在 PDF 第 10 頁）。

  python -X utf8 scripts/god_evolution_read_aloud.py --check
  python -X utf8 scripts/god_evolution_read_aloud.py --out "G:/我的雲端硬碟/神的演化/神的演化（分章朗讀）"
"""
from __future__ import annotations

import argparse
import glob
import json
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import scan_ocr  # noqa: E402

OFFSET = 6
LAST_PDF = 526
# (檔名序號, 標題, 印刷起始頁)
TOC = [
    ("00", "序言", 4),
    ("01", "第01章　基原信仰", 10), ("02", "第02章　薩滿師", 31),
    ("03", "第03章　酋邦時代的宗教", 48), ("04", "第04章　古代國家的眾神", 72),
    ("05", "第05章　古以色列人的宗教：多神教", 100), ("06", "第06章　從多神到一神崇拜", 139),
    ("07", "第07章　從一神崇拜到一神教", 180), ("08", "第08章　斐洛的故事", 205),
    ("09", "第09章　邏各斯：神的運算法則", 236),
    ("10", "第10章　耶穌行過些什麼事？", 266), ("11", "第11章　愛的使徒", 288),
    ("12", "第12章　最適基督教生存", 313), ("13", "第13章　耶穌如何會變成救主？", 329),
    ("14", "第14章　《古蘭經》", 356), ("15", "第15章　麥加", 371), ("16", "第16章　麥地那", 382),
    ("17", "第17章　聖戰", 404), ("18", "第18章　穆罕默德", 420),
    ("19", "第19章　道德想像力", 440), ("20", "第20章　我們是獨一無二的嗎？", 460),
    ("21", "跋　順道一談：上帝是什麼樣子的？", 472),
    ("22", "附錄　宗教如何源自於人類天性", 487),
    ("23", "經文譯本小識", 510), ("24", "鳴謝", 516),
]
# 各部的扉頁（只有部名與題詞）不進任何一章
PART_PAGES = {8, 9, 98, 99, 264, 265, 354, 355, 438, 439}
_END = tuple("。！？」』）…—")
_PUNCT = re.compile(r"[，。、；：？！「」『』（）]")


from chaohwei_build import to_fullwidth_punct  # noqa: E402


def unwrap(text: str) -> str:
    """Gemini 在這本照印刷行斷句：一行約 40 字。上一行印滿（≥ 常見行長 −2）就是同一段。"""
    lines = text.split("\n")
    body = [len(l) for l in lines if l.strip() and not l.startswith("[^")]
    if len(body) < 5:
        return text
    full = max(set(body), key=body.count)
    if full < 20 or sum(1 for n in body if n >= full - 2) < len(body) * 0.5:
        return text
    out: list[str] = []
    last = 0
    for ln in lines:
        if out and not ln.startswith("[^") and not out[-1].startswith("[^") and ln.strip() and last >= full - 2:
            out[-1] += ln
        else:
            out.append(ln)
        last = len(ln)
    return "\n".join(out)


def load_pages() -> dict[int, dict]:
    """PDF 頁 → {header, body}。Gemini 優先，缺的用 MinerU。"""
    pages: dict[int, dict] = {}
    for f in glob.glob("c:/tmp/evo/part_*.jsonl"):
        for ln in open(f, encoding="utf-8"):
            if not ln.startswith("{"):
                continue
            c = json.loads(ln)
            body = c["content"].split("———————————————")[0]   # 註腳區不要
            # MinerU 一個 block 才是一段，block 內的換行是直排的一行；接成一段一行
            body = "\n".join(b.replace("\n", "") for b in body.split("\n\n"))
            pages[c["page_number"]] = {"header": " ".join(c.get("headers") or []),
                                       "body": body, "engine": "mineru"}
    for f in glob.glob("c:/tmp/evo/gem/*.json"):
        d = json.loads(open(f, encoding="utf-8").read())
        s, e = d["start"], d["end"]
        # 🚨 不信模型回報的 page：它常填印刷頁碼，被 ocr_pdf 換算到隔壁幾頁去
        #    （序言的段落跑進第一章）。頁數剛好就照回傳順序對回，不對就整批不用、交給 MinerU。
        if len(d["pages"]) != e - s + 1:
            continue
        got = [{**p, "page": s + i} for i, p in enumerate(d["pages"])]
        for p in got:
            printed, header, body = scan_ocr.split_printed((p.get("text") or "").replace("\\n", "\n"))
            body = to_fullwidth_punct(unwrap(body))
            if body.strip() and body.strip() != "（無正文）":
                pages[int(p["page"])] = {"header": header, "body": body, "engine": "gemini"}
    return pages


def page_paras(body: str) -> list[str]:
    out = []
    for ln in body.split("\n"):
        ln = ln.strip()
        if not ln or ln.startswith("[^") or ln.startswith("<table") or re.match(r"^\d{1,3}$", ln):
            continue
        ln = re.sub(r"\[\^[^\]]*\]", "", ln)           # 正文註號
        out.append(ln)
    return out


def is_heading(p: str) -> bool:
    return len(p) <= 22 and not _PUNCT.search(p[:-1] if p.endswith("？") else p) and not p.isascii()


def build(pages: dict[int, dict]):
    starts = [(no, title, pr + OFFSET) for no, title, pr in TOC]
    chapters = []
    for i, (no, title, s) in enumerate(starts):
        e = (starts[i + 1][2] - 1) if i + 1 < len(starts) else LAST_PDF
        paras: list[str] = []
        missing = []
        for pdf in range(s, e + 1):
            if pdf - OFFSET in PART_PAGES:
                continue
            pg = pages.get(pdf)
            if not pg:
                missing.append(pdf)
                continue
            ps = page_paras(pg["body"])
            if pdf == s:     # 章首頁：丟掉章名、英文章名、「第N章」
                core = re.sub(r"^第\d+章　", "", title).split("　")[-1]
                ps = [p for p in ps if not (p.isascii() or re.fullmatch(r"第\s*\d+\s*章", p)
                                            or p.replace(" ", "") in (core, title.replace("　", "")))]
            for j, p in enumerate(ps):
                if j == 0 and paras and not paras[-1].endswith(_END) and not is_heading(paras[-1]) \
                        and not is_heading(p):
                    paras[-1] += p          # 跨頁續段
                else:
                    paras.append(p)
        chapters.append({"no": no, "title": title, "paras": paras, "missing": missing})
    return chapters


def write_docx(chapters, out: Path) -> None:
    import docx
    out.mkdir(parents=True, exist_ok=True)
    for ch in chapters:
        d = docx.Document()
        d.add_heading(ch["title"], level=1)
        for p in ch["paras"]:
            if is_heading(p):
                d.add_heading(p, level=2)
            else:
                d.add_paragraph(p)
        name = f"{ch['no']} {ch['title'].replace('　', ' ')}.docx"
        d.save(str(out / re.sub(r'[\\/:*?"<>|]', "", name)))


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--check", action="store_true")
    ap.add_argument("--out")
    a = ap.parse_args()
    pages = load_pages()
    eng = {k: v["engine"] for k, v in pages.items()}
    print(f"有內容的頁 {len(pages)}（gemini {sum(1 for v in eng.values() if v == 'gemini')}、"
          f"mineru {sum(1 for v in eng.values() if v == 'mineru')}）")
    chs = build(pages)
    for c in chs:
        n = sum(len(p) for p in c["paras"])
        print(f"  {c['no']} {c['title'][:22]:<24} {len(c['paras']):>4} 段 {n:>7,} 字"
              + (f"  缺頁 {c['missing']}" if c["missing"] else ""))
    if a.out:
        write_docx(chs, Path(a.out))
        print(f"✓ 寫出 {len(chs)} 檔 → {a.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
