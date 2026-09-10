# -*- coding: utf-8 -*-
"""把全集某一卷的**繁中譯文**出成 Word 檔（純中文，供朗讀軟體讀）。

reader 是雙欄對照（原文＋繁中），朗讀用不需要原文——這支只取 zh 欄，按章編排成
一份乾淨的 .docx：書名頁、每章一個 Heading 1、引文段落縮排（`> ` 標記剝掉，
不然朗讀軟體會把大於號念出來）。

資料來源是各作者模組的 checkpoint（`<author>_data/<slug>/secN.json`），跟
uchimura_auto.py 共用同一批檔案，所以翻到哪裡就出到哪裡；沒有譯文的段落退回原文。

  python scripts/collected_works_docx.py --author howes --work howes-prophet \
      --out "C:/Users/user/Desktop/日本的現代先知_內村鑑三_中譯.docx"
"""
from __future__ import annotations

import argparse
import importlib
import json
import sys
from pathlib import Path

SCRIPT_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPT_DIR))

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

import uchimura_auto as ua  # noqa: E402  (AUTHOR_MODULES + 同一套 checkpoint 路徑)

BODY_FONT = "Microsoft JhengHei"  # 微軟正黑體：繁中字面完整，朗讀軟體不挑字型


def _set_cjk(run, font: str = BODY_FONT) -> None:
    """python-docx 只設 ascii 字型；中文要另外寫 w:eastAsia，否則 Word 會退回細明體。"""
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)


def load_sections(author: str, slug: str) -> tuple[dict, list[dict]]:
    mod = importlib.import_module(ua.AUTHOR_MODULES[author])
    data_root = SCRIPT_DIR.parent / ".claude" / "skills" / "ebook-collected-works" / \
        getattr(mod, "DATA_DIRNAME", "uchimura_data")
    out = []
    for i, sec in enumerate(mod.load_work_sections(slug)):
        cp = data_root / slug / f"sec{i}.json"
        cache = json.loads(cp.read_text(encoding="utf-8")) if cp.exists() else {}
        src = list(sec["paras"])
        zh = (list(cache.get("zh") or []) + [None] * len(src))[:len(src)]
        out.append({
            "title": cache.get("title_zh") or sec.get("title_zh") or sec["heading"],
            # 沒譯到的段落退回原文，寧可中英夾雜也不要缺內容
            "paras": [(z or s) for z, s in zip(zh, src)],
            # 原書印刷頁碼（沒有這個欄位的作者模組就整段留 None）
            "pages": list(sec.get("pages") or [None] * len(src))[:len(src)],
        })
    return mod.REGISTRY[slug], out


def build(author: str, slug: str, out_path: Path,
          page_marks: bool = False) -> tuple[int, int]:
    mod = importlib.import_module(ua.AUTHOR_MODULES[author])
    work, sections = load_sections(author, slug)

    doc = Document()
    st = doc.styles["Normal"]
    st.font.size = Pt(12)
    st.font.name = BODY_FONT
    st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    st.paragraph_format.line_spacing = 1.5
    st.paragraph_format.space_after = Pt(6)

    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(2.2)
        s.left_margin = s.right_margin = Cm(2.4)

    # ── 書名頁 ──
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(work["title"])
    r.bold = True
    r.font.size = Pt(24)
    _set_cjk(r)

    for line, size in [(work.get("original_title", ""), 12),
                       (f"{getattr(mod, 'AUTHOR_ZH', '')}（{getattr(mod, 'AUTHOR_EN', '')}）著", 13),
                       (f"{work.get('year', '')}", 11)]:
        if not line:
            continue
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = p.add_run(str(line))
        rr.font.size = Pt(size)
        _set_cjk(rr)

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    nr = note.add_run("繁體中文譯稿（逐段機器翻譯，供朗讀使用；引文以縮排表示）")
    nr.font.size = Pt(10)
    nr.italic = True
    _set_cjk(nr)
    doc.add_page_break()

    n_par = 0
    last_page = None
    for sec in sections:
        h = doc.add_heading(level=1)
        hr = h.add_run(sec["title"])
        hr.font.size = Pt(17)
        _set_cjk(hr)
        last_page = None
        for para, page in zip(sec["paras"], sec["pages"]):
            # 🚨 頁碼標記預設關閉：這份 docx 的用途是朗讀，朗讀軟體會把「頁 21」
            # 唸出來。要引用的人才開 --page-marks，另出一份。
            if page_marks and page and page != last_page:
                m = doc.add_paragraph()
                m.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                m.paragraph_format.space_after = Pt(0)
                mr = m.add_run(f"〔原書 p. {page}〕")
                mr.font.size = Pt(8)
                mr.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
                _set_cjk(mr)
            last_page = page or last_page
            text = (para or "").strip()
            if not text:
                continue
            quoted = text.startswith("> ")
            if quoted:
                text = text[2:].strip()
            p = doc.add_paragraph()
            if quoted:
                p.paragraph_format.left_indent = Cm(1.2)
                p.paragraph_format.right_indent = Cm(0.6)
            pr = p.add_run(text)
            pr.font.size = Pt(12)
            _set_cjk(pr)
            n_par += 1

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return len(sections), n_par


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--author", required=True, choices=sorted(ua.AUTHOR_MODULES))
    ap.add_argument("--work", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--page-marks", action="store_true",
                    help="每逢原書換頁插一個〔原書 p. N〕標記（供引註；朗讀用不要開）")
    args = ap.parse_args()
    n_sec, n_par = build(args.author, args.work, Path(args.out), page_marks=args.page_marks)
    print(f"寫出 {args.out}\n  {n_sec} 章 / {n_par} 段")


if __name__ == "__main__":
    main()
