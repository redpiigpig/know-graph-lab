# -*- coding: utf-8 -*-
"""備好五個章則頁要附在頁尾的 Word 檔：用最新版內容，但**改成正式的文件名稱**
（去掉「英網-N」這個編輯室內部批號與日期尾碼），並把 AI 規範裡的「學生」改成「作者」。

🚨 改 docx 的字要逐個 run 改：同一段會被 Word 切成多個 run，只改 paragraph.text
   會把整段的格式（字體、粗體）壓平。
🚨 「學生」只出現在那一句，改完要印出改了幾處；0 處就是沒改到，不能當成成功。

用法：python -X utf8 scripts/hcjbs_cms_docx.py [--out c:/tmp/hcu-cms/docx]
"""
import argparse
import shutil
from pathlib import Path

import docx

SRC = Path('stores/玄奘佛學研究')
# 正本檔名 → 對外要用的名稱
RENAME = {
    '英網-1 投稿指引(徵稿函).docx': '投稿指引（徵稿函）.docx',
    '英網-2 編輯團隊資訊.docx': '編輯團隊資訊.docx',
    '英網-3 草擬學術倫理聲明260425.docx': '學術倫理聲明.docx',
    '英網-4 學報AI使用規範260904.docx': '學報AI使用規範.docx',
    '英網-5 期刊審查流程260213.docx': '期刊審查流程.docx',
}
# 依老師指示：AI 使用規範的「學生」改「作者」
REPLACE = {'學報AI使用規範.docx': [('學生使用相關工具', '作者使用相關工具')]}


def patch(path, pairs):
    d = docx.Document(str(path))
    hits = 0
    for para in d.paragraphs:
        for run in para.runs:
            for a, b in pairs:
                if a in run.text:
                    run.text = run.text.replace(a, b)
                    hits += 1
    # 表格裡也可能有
    for tbl in d.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        for a, b in pairs:
                            if a in run.text:
                                run.text = run.text.replace(a, b)
                                hits += 1
    if hits:
        d.save(str(path))
    return hits


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--out', default='c:/tmp/hcu-cms/docx')
    args = ap.parse_args()
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)

    bad = 0
    for src_name, dst_name in RENAME.items():
        src = SRC / src_name
        if not src.exists():
            print(f'❌ 找不到正本：{src}')
            bad += 1
            continue
        dst = out / dst_name
        shutil.copy(src, dst)
        note = ''
        if dst_name in REPLACE:
            hits = patch(dst, REPLACE[dst_name])
            note = f'  「學生」→「作者」改了 {hits} 處'
            if not hits:
                print(f'❌ {dst_name} 一處都沒改到——檢查字串是不是被 Word 切成多個 run')
                bad += 1
        print(f'✔ {dst_name}  {dst.stat().st_size // 1024} KB{note}')

    # 驗收：把改過那份重新讀一遍，確認文字真的是「作者」
    ai = out / '學報AI使用規範.docx'
    if ai.exists():
        text = '\n'.join(p.text for p in docx.Document(str(ai)).paragraphs)
        ok = '作者使用相關工具' in text and '學生使用相關工具' not in text
        print(f'{"✅" if ok else "❌"} 驗收 AI 規範：作者={("作者使用相關工具" in text)} 學生殘留={("學生使用相關工具" in text)}')
        if not ok:
            bad += 1
    print(f'\n{"✅ 全部備妥" if not bad else f"❌ {bad} 項有問題"}  → {out}')
    return 1 if bad else 0


if __name__ == '__main__':
    raise SystemExit(main())
