# -*- coding: utf-8 -*-
"""全集每一本書 → Drive `全集/{學科}/{作者}/{書名}.docx`（繁中讀本）。

使用者 2026-09-23：「Drive 也都要有一本 Word」。網站讀的是 JSONL（Drive `_chunks/`＋
R2 副本），Drive 的 `全集/` 夾原本只放手動下載的原始檔，所以大多數全集在 Drive 上
沒有一份能直接打開來讀的東西。這支把**每一本**全集書從 `_chunks/{id}.jsonl` 出成 Word。

與 `collected_works_docx.py` 的分工：那支讀 uchimura_auto 系列的 checkpoint、專供朗讀
（純中文、可分章），只接得上少數作者；這支讀所有全集共同的 JSONL，所以印順、星雲、
柏拉圖、阿奎那……都出得來。

排版：
  * 標題頁（書名／作者／原題）→ 章節標題取自 chapter_path，逐層做成 Heading 1–3
    （Word 導覽窗格可跳）。
  * 正文只放繁中（content）。原文對照在網站上看；這是讀本不是對照本。
  * 引用號：chunk 帶 anchors 且與段落一一對應時（Stephanus 17a、阿奎那 I q1 a1），
    段首灰字標出；否則 page_number 變動時在頁首標「〔頁 46〕」。沒有就不標，不捏造。
  * markdown：`## ` 小標、`> ` 引文縮排、`**粗體**`、表格畫成真表格、`[^n]: ` 註文小字。

增量：JSONL 比 Word 新（或 Word 不存在）才重出，所以翻譯還在進行的書重跑就會更新
（[[feedback_deliverable_all_copies]]）。

  node scripts/dump_collected_works_authors.mjs > output/collected_works_authors.json
  python -X utf8 scripts/collected_works_word.py --dry           # 只列會出哪些
  python -X utf8 scripts/collected_works_word.py --only <ebook_id> [...]
  python -X utf8 scripts/collected_works_word.py                 # 全部（增量）
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys
import time
from pathlib import Path

SCRIPT_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPT_DIR))
sys.stdout.reconfigure(encoding="utf-8")

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

from collected_works_docx import _add_table, _set_cjk  # noqa: E402

REPO = SCRIPT_DIR.parent
DRIVE_CW = Path(r"G:\我的雲端硬碟\資料\知識圖工作室\全集")
AUTHORS_JSON = REPO / "output" / "collected_works_authors.json"
MANIFEST = REPO / "output" / "collected_works_word_manifest.json"
GRAY = RGBColor(0x88, 0x88, 0x88)


def safe_name(s: str, limit: int = 100) -> str:
    s = re.sub(r'[\\/:*?"<>|\r\n\t]', "_", s).strip(" .")
    return s[:limit].rstrip(" .") or "untitled"


def _clean_inline(t: str) -> str:
    t = re.sub(r"\*\*(.+?)\*\*", r"\1", t)
    t = re.sub(r"(?<!\*)\*(?!\s)(.+?)(?<!\s)\*(?!\*)", r"\1", t)
    return t


def _body(doc, text: str, *, indent: bool = False, small: bool = False, mark: str | None = None):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.left_indent = Cm(1.0)
    if mark:
        r = p.add_run(f"〔{mark}〕")
        r.font.size = Pt(8.5)
        r.font.color.rgb = GRAY
        _set_cjk(r)
    lines = text.split("\n")
    for k, line in enumerate(lines):
        r = p.add_run(_clean_inline(line))
        r.font.size = Pt(9.5 if small else 11.5)
        _set_cjk(r)
        if k < len(lines) - 1:
            r.add_break()
    return p


def _level(t: str) -> str:
    """章節路徑的一層 → 標題文字。

    * 去掉上架時切分長章留下的「（2）」：那是 reader 分頁的編號不是章節，不去掉的話
      豪斯評傳的〈導論〉會變成〈導論（1）〉〈導論（2）〉〈導論（3）〉三個標題。
    * 去掉星雲全集路徑裡的頁碼前綴「p046　」（頁碼另外以〔頁 46〕標出）。"""
    t = t.strip()
    t = re.sub(r"（\d+）$", "", t)
    t = re.sub(r"^p\d+[\s　]+", "", t)
    return t.strip()


def book_title(t: str) -> str:
    """資料庫書名常帶「（希英繁三欄）」「（英繁對照）」這類版本註記；Word 只放繁中，
    照抄會誤導，去掉。"""
    return re.sub(r"\s*（[^（）]*(對照|三欄|雙語|雙欄|原文)[^（）]*）\s*$", "", t).strip() or t


def _heading(doc, text: str, level: int):
    h = doc.add_heading(level=min(level, 3))
    r = h.add_run(_clean_inline(text))
    _set_cjk(r)


def render_chunk(doc, c: dict, state: dict) -> None:
    path = [_level(x) for x in (c.get("chapter_path") or "").split(" · ") if x.strip()]
    # 第一層通常是書名本身（標題頁已經有了），從第二層起做章節標題
    levels = [x for x in (path[1:] if len(path) > 1 else path) if x]
    prev = state.get("path", [])
    same = True
    for i, t in enumerate(levels):
        same = same and i < len(prev) and prev[i] == t
        if not same:                        # 從第一個不同的層級起，往下都要印標題
            _heading(doc, t, i + 1)
            state["last_head"] = t
    state["path"] = levels

    paras = [p for p in re.split(r"\n\s*\n", c.get("content") or "") if p.strip()]
    anchors = c.get("anchors") or []
    # anchors 與「正文段落」一一對應（小標列不算）；對不上就改用頁碼
    body_idx = [i for i, p in enumerate(paras) if not p.lstrip().startswith("#")]
    use_anchor = bool(anchors) and len(anchors) >= len(body_idx) and any(anchors)
    page = c.get("page_number")
    if not use_anchor and isinstance(page, int) and page > 0 and page != state.get("page"):
        state["page"] = page
        state["pending_page"] = f"頁 {page}"
    k = 0
    for p in paras:
        s = p.strip()
        if s.startswith("#"):
            t = _level(s.lstrip("#").strip())
            # 正文開頭的小標常只是 chapter_path 的重述（阿奎那「第1題 … · 第1節」），跳過
            dup = t in levels or any(t == " · ".join(levels[-k:]) for k in range(1, len(levels) + 1))
            if t and not dup and t != state.get("last_head"):
                _heading(doc, t, min(len(levels) + 1, 3))
                state["last_head"] = t
            continue
        if s.startswith("|") and "\n|" in s:
            _add_table(doc, s)
            k += 1
            continue
        mark = None
        if use_anchor:
            mark = anchors[k] if k < len(anchors) and anchors[k] else None
        elif state.get("pending_page"):
            mark = state.pop("pending_page")
        if s.startswith(">"):
            _body(doc, re.sub(r"(?m)^>\s?", "", s), indent=True, mark=mark)
        elif re.match(r"\[\^\w+\]:", s) or s.startswith("—" * 5):
            _body(doc, s, small=True, mark=mark)
        else:
            _body(doc, s, mark=mark)
        k += 1


def build_docx(meta: dict, chunks: list[dict], out: Path) -> int:
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Microsoft JhengHei"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(2.2)
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(meta["title"])
    r.bold = True
    r.font.size = Pt(22)
    _set_cjk(r)
    for line, size in ((meta.get("author", ""), 14), (meta.get("original", ""), 12), (meta.get("subtitle", ""), 10)):
        if line:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rr = p.add_run(line)
            rr.font.size = Pt(size)
            _set_cjk(rr)
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rn = note.add_run(f"知識圖工作室‧全集讀本（由站上全文匯出，{time.strftime('%Y-%m-%d')}）")
    rn.font.size = Pt(9)
    rn.font.color.rgb = GRAY
    _set_cjk(rn)
    doc.add_page_break()
    state: dict = {}
    n = 0
    for c in chunks:
        if c.get("chunk_type") == "cover":
            continue
        render_chunk(doc, c, state)
        n += 1
    out.parent.mkdir(parents=True, exist_ok=True)
    tmp = out.with_suffix(".docx.part")
    doc.save(tmp)
    os.replace(tmp, out)
    return n


def _plain(s: str) -> str:
    return re.sub(r"[‧・·\s.]", "", s or "")


def author_folder(disc: str, who: str) -> Path:
    """Drive 上已經有這位作者的資料夾就沿用（原始檔先前是用短名建的：「伊利亞德」
    而不是「米爾恰‧伊利亞德」），否則用全名新建。"""
    base = DRIVE_CW / safe_name(disc, 30)
    if base.exists():
        for d in base.iterdir():
            a, b = _plain(d.name), _plain(who)
            if d.is_dir() and min(len(a), len(b)) >= 2 and (a.endswith(b) or b.endswith(a)):
                return d
    return base / safe_name(who, 40)


def plan(authors: list[dict], rows: list[dict], chunks_dir: str) -> list[dict]:
    by_id: dict[str, tuple[dict, dict]] = {}
    by_name: dict[str, dict] = {}
    for a in authors:
        for key in (a["name"], a.get("nameEn"), a.get("nameOriginal")):
            if key:
                by_name.setdefault(_plain(key).lower(), a)
        for w in a["works"]:
            if w["ebookId"]:
                by_id.setdefault(w["ebookId"], (a, w))
    muller = next((a for a in authors if a["slug"] in ("max-mueller", "mueller")), None)
    jobs, used = [], set()
    for row in rows:
        src = Path(chunks_dir) / f"{row['id']}.jsonl"
        if not src.exists():
            continue
        sub = ""
        title, original = book_title(row.get("title") or row["id"]), row.get("original_title") or ""
        if row["id"] in by_id:
            a, w = by_id[row["id"]]
            disc, who = a["disciplineGroup"], a["name"]
            original = w.get("titleOriginal") or original
        elif row.get("category") == "世界宗教" and muller:
            # 《東方聖書》各卷：hub 以 externalUrl 連到 /sacred-books-east，沒有逐卷 ebookId。
            # 放在主編穆勒底下，而不是按各卷譯者開一堆資料夾。
            disc, who, sub = muller["disciplineGroup"], muller["name"], "東方聖書"
        else:
            a = by_name.get(_plain(row.get("author") or "").lower())
            disc = a["disciplineGroup"] if a else (row.get("category") or "其他")
            who = a["name"] if a else (row.get("author") or "佚名")
        folder = author_folder(disc, who)
        if sub:
            folder = folder / sub
        name = safe_name(title)
        out = folder / f"{name}.docx"
        i = 2
        while str(out).lower() in used:
            out = folder / f"{name}（{i}）.docx"
            i += 1
        used.add(str(out).lower())
        jobs.append({"id": row["id"], "src": str(src), "out": str(out), "title": title,
                     "author": who, "discipline": disc, "original": original,
                     "subtitle": row.get("subtitle") or ""})
    return jobs


def fetch_rows() -> list[dict]:
    import requests
    import translate_ebook_to_zh as te
    rows, off = [], 0
    while True:
        r = requests.get(f"{te.URL}/rest/v1/ebooks", headers={**te.H_JSON, "Range": f"{off}-{off + 999}"},
                         params={"select": "id,title,subtitle,author,original_title,category",
                                 "collection": "eq.collected-works", "order": "id"}, timeout=60)
        r.raise_for_status()
        b = r.json()
        rows += b
        if len(b) < 1000:
            return rows
        off += 1000


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry", action="store_true")
    ap.add_argument("--only", nargs="*")
    ap.add_argument("--force", action="store_true", help="不管新舊一律重出")
    args = ap.parse_args()
    import translate_ebook_to_zh as te
    # 同時只准一支：fleet keeper 每 30 分會拉一次，手動也可能在跑，兩支寫同一批檔會互蓋
    lock = REPO / "output" / ".collected_works_word.lock"
    if lock.exists():
        try:
            other = int(lock.read_text().strip() or 0)
        except ValueError:
            other = 0
        if other and _pid_alive(other):
            print("另一支正在匯出，這次略過", flush=True)
            return
    lock.parent.mkdir(parents=True, exist_ok=True)
    lock.write_text(str(os.getpid()))
    try:
        _run(args, te)
    finally:
        lock.unlink(missing_ok=True)


def _pid_alive(pid: int) -> bool:
    """Windows：OpenProcess＋GetExitCodeProcess。🚨 不可用 os.kill(pid, 0)——
    在 Windows 上那會呼叫 TerminateProcess，等於把對方殺掉。"""
    import ctypes
    k = ctypes.windll.kernel32
    h = k.OpenProcess(0x1000, False, pid)          # PROCESS_QUERY_LIMITED_INFORMATION
    if not h:
        return False
    code = ctypes.c_ulong()
    ok = k.GetExitCodeProcess(h, ctypes.byref(code))
    k.CloseHandle(h)
    return bool(ok) and code.value == 259          # STILL_ACTIVE


def _run(args, te) -> None:
    if not DRIVE_CW.exists():
        sys.exit(f"找不到 {DRIVE_CW}（G: 沒掛？先 Test-Path 'G:\\我的雲端硬碟'）")
    store = REPO / "stores" / "collectedWorks.ts"
    if not AUTHORS_JSON.exists() or AUTHORS_JSON.stat().st_mtime < store.stat().st_mtime:
        # 全集頁面改過（新作者、新卷）就重倒作者清單，否則新書會被歸到「其他」
        import subprocess
        AUTHORS_JSON.parent.mkdir(parents=True, exist_ok=True)
        dump = subprocess.run(["node", str(SCRIPT_DIR / "dump_collected_works_authors.mjs")],
                              cwd=REPO, capture_output=True, text=True, encoding="utf-8")
        if dump.returncode != 0 or not dump.stdout.strip():
            sys.exit(f"作者清單匯出失敗：{dump.stderr[-300:]}")
        AUTHORS_JSON.write_text(dump.stdout, encoding="utf-8")
    authors = json.loads(AUTHORS_JSON.read_text(encoding="utf-8"))
    jobs = plan(authors, fetch_rows(), te.CHUNKS_DIR)
    if args.only:
        jobs = [j for j in jobs if j["id"] in set(args.only)]
    print(f"全集書 {len(jobs)} 本有 JSONL", flush=True)
    done = skipped = failed = 0
    manifest = []
    for j in jobs:
        out, src = Path(j["out"]), Path(j["src"])
        fresh = out.exists() and out.stat().st_mtime >= src.stat().st_mtime
        manifest.append(j)
        if args.dry:
            print(("  ✓ " if fresh else "  + ") + str(out.relative_to(DRIVE_CW)))
            continue
        if fresh and not args.force:
            skipped += 1
            continue
        try:
            chunks = [json.loads(l) for l in src.read_text(encoding="utf-8").splitlines() if l.strip()]
            n = build_docx(j, chunks, out)
            done += 1
            print(f"  [{done + skipped}/{len(jobs)}] {out.relative_to(DRIVE_CW)}（{n} 節）", flush=True)
        except Exception as e:  # noqa: BLE001 — 一本壞了不擋其他本
            failed += 1
            print(f"  ⚠ {j['title'][:30]}：{str(e)[:160]}", flush=True)
    MANIFEST.write_text(json.dumps(manifest, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"完成：新出 {done}、沿用 {skipped}、失敗 {failed}", flush=True)


if __name__ == "__main__":
    main()
